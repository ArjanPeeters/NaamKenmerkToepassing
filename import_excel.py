"""Excel -> NaamKenmerkToepassing.db importer.

Vult naam / kenmerk / toepassing + de koppeltabellen en select_nlsfb bij vanuit
de "EenduidigeMateriaalbenaming" .xlsm. extralijsten en select_ral staan NIET in
de Excel en worden met rust gelaten.

Gebruik:  python import_excel.py ["pad\\naar\\lijst vX.xlsm"]

Bronnen in de Excel:
  sheet 'NAA.K.T.'      rij 7+ : B=naam C=kenmerk E=toepassing(voorbeeld) G+H=NL-SfB
  sheet 'TOEPASSINGEN'         : kolom per naam (header), eronder de toepassingen
"""
import shutil
import sys
from datetime import datetime

import openpyxl
from flask import Flask

from dbModels import db, Naam, Kenmerk, Toepassing, Select_NLSFB, Synoniem

DEFAULT_XLSX = r'C:\OneDrive\OneDrive - BIMnerd\BIM repo\NAAKT\EenduidigeMateriaalbenaming_lijst v2.5.xlsm'
DB_PATH = 'instance/NaamKenmerkToepassing.db'


def norm(v):
    return str(v).strip().lower() if v is not None and str(v).strip() else None


def parse_excel(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)

    # --- NAA.K.T.: namen, kenmerken per naam, NL-SfB + synoniemen per naam_kenmerk ---
    kenmerken = {}   # naam -> [kenmerk, ...] (volgorde behouden, uniek)
    nlsfb = {}       # 'naam_kenmerk' -> 'f2'
    synoniemen = {}  # 'naam_kenmerk' -> 'zoekwoorden ...' (kolom D)
    for row in wb['NAA.K.T.'].iter_rows(min_row=7, values_only=True):
        naam, kenmerk = norm(row[1]), norm(row[2])
        if not naam:
            continue
        kenmerken.setdefault(naam, [])
        if kenmerk:
            if kenmerk not in kenmerken[naam]:
                kenmerken[naam].append(kenmerk)
            code = (str(row[6]).strip() if row[6] else '') + (str(row[7]).strip() if row[7] else '')
            if code:
                nlsfb.setdefault(f'{naam}_{kenmerk}', code)
            syn = str(row[3]).strip().lower() if row[3] else ''
            if syn:
                synoniemen.setdefault(f'{naam}_{kenmerk}', syn)

    # --- TOEPASSINGEN: kolom per naam ---
    toepassingen = {}  # naam -> [toepassing, ...]
    rows = list(wb['TOEPASSINGEN'].iter_rows(values_only=True))
    naam_col = {c: norm(h) for c, h in enumerate(rows[0])
                if h and not str(h).startswith('Kolom')}
    for c, naam in naam_col.items():
        toepassingen.setdefault(naam, [])
    for row in rows[1:]:
        for c, naam in naam_col.items():
            t = norm(row[c]) if c < len(row) else None
            if t and t not in toepassingen[naam]:
                toepassingen[naam].append(t)

    return kenmerken, toepassingen, nlsfb, synoniemen


def get_or_create(model, attr, value, cache):
    if value in cache:
        return cache[value]
    obj = model.query.filter(getattr(model, attr) == value).first()
    if obj is None:
        obj = model(**{attr: value})
        db.session.add(obj)
    cache[value] = obj
    return obj


def run(xlsx_path):
    shutil.copy2(DB_PATH, f'{DB_PATH}.bak')
    print(f'Backup -> {DB_PATH}.bak')

    kenmerken, toepassingen, nlsfb, synoniemen = parse_excel(xlsx_path)
    alle_namen = set(kenmerken) | set(toepassingen)

    app = Flask(__name__)
    app.config['SQLALCHEMY_DATABASE_URI'] = f'sqlite:///{DB_PATH.split("/")[-1]}'
    app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
    db.init_app(app)

    with app.app_context():
        db.create_all()  # maakt ontbrekende tabellen (o.a. synoniem) aan, laat bestaande met rust
        k_cache, t_cache = {}, {}
        for naam in sorted(alle_namen):
            n = get_or_create(Naam, 'naam', naam, {})
            n.kenmerken = [get_or_create(Kenmerk, 'kenmerk', k, k_cache)
                           for k in kenmerken.get(naam, [])]
            n.toepassingen = [get_or_create(Toepassing, 'toepassing', t, t_cache)
                              for t in toepassingen.get(naam, [])]

        for materiaal, code in nlsfb.items():
            obj = Select_NLSFB.query.filter_by(materiaal=materiaal).first()
            if obj is None:
                db.session.add(Select_NLSFB(materiaal=materiaal, nlsfb=code))
            else:
                obj.nlsfb = code

        for materiaal, woorden in synoniemen.items():
            obj = Synoniem.query.filter_by(materiaal=materiaal).first()
            if obj is None:
                db.session.add(Synoniem(materiaal=materiaal, woorden=woorden))
            else:
                obj.woorden = woorden

        db.session.flush()
        # wees-records opruimen (kenmerk/toepassing zonder naam)
        for obj in Kenmerk.query.all():
            if not obj.namen:
                db.session.delete(obj)
        for obj in Toepassing.query.all():
            if not obj.namen:
                db.session.delete(obj)
        db.session.commit()

        print(f'namen={Naam.query.count()} kenmerken={Kenmerk.query.count()} '
              f'toepassingen={Toepassing.query.count()} nlsfb={Select_NLSFB.query.count()} '
              f'synoniemen={Synoniem.query.count()}')


if __name__ == '__main__':
    run(sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX)
